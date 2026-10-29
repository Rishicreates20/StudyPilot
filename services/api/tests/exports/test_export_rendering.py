"""Rendering exports: valid PDF and PPTX files, readable text, embedded fonts, honest Unicode.

These tests check the actual output: files are re-opened with independent readers (pypdf for PDF
structure and text, PDFium for drawing pages, python-pptx and zipfile for PowerPoint). Text
extraction of shaped Indic text is unreliable in every PDF reader, so Hindi and Odia are checked
by font coverage, embedded fonts and a drawn page, not by extracting the text back.
"""

import io
import re
import zipfile
from typing import Any

import pypdfium2  # pyright: ignore[reportMissingTypeStubs]
import pytest
from pptx import Presentation
from pypdf import PdfReader

from app.ai.lesson.schema import Lesson
from app.exports.lesson_pdf import _checklist, _info  # pyright: ignore[reportPrivateUsage]
from app.exports.model import PDF_TYPE, PPTX_TYPE, RenderedFile, file_name_for
from app.exports.pdf_base import BODY_FAMILY, StudyDocument
from app.exports.render import UnsupportedExport, render
from app.exports.slides import (
    MAX_CODE_LINES,
    MAX_POINT_CHARS,
    MAX_POINTS,
    first_sentence,
    lesson_outline,
    plain,
    roadmap_outline,
    shorten,
)
from app.exports.typography import REPLACEMENT, TextGuard, coverage
from support.exports import (
    HINDI_TEXT,
    ODIA_TEXT,
    cited_sources,
    indic_lesson,
    lesson_source,
    roadmap_source,
)
from support.lessons import valid_lesson

FILE_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")  # the database's rule


def pdf_text(result: RenderedFile) -> str:
    reader = PdfReader(io.BytesIO(result.data))
    return "\n".join(page.extract_text() for page in reader.pages)


def embedded_fonts(result: RenderedFile) -> set[str]:
    reader = PdfReader(io.BytesIO(result.data))
    names: set[str] = set()
    for page in reader.pages:
        resources: Any = page["/Resources"]
        fonts: Any = resources.get("/Font", {})
        for font in fonts.values():
            names.add(str(font.get_object()["/BaseFont"]).split("+", 1)[-1])
    return names


def lesson_with(edit: Any) -> Lesson:
    data = valid_lesson().model_dump()
    edit(data)
    return Lesson.model_validate(data)


# --- PDF ----------------------------------------------------------------------------------------


@pytest.mark.parametrize("export_format", ["notes_pdf", "revision_pdf"])
def test_a_lesson_pdf_is_a_valid_pdf_with_its_content(export_format: str) -> None:
    source = lesson_source(citations=cited_sources())
    result = render(export_format, source)  # type: ignore[arg-type]

    assert result.content_type == PDF_TYPE
    assert result.data.startswith(b"%PDF-") and result.data.rstrip().endswith(b"%%EOF")
    reader = PdfReader(io.BytesIO(result.data))
    assert len(reader.pages) == result.page_count >= 1
    assert reader.metadata is not None and reader.metadata.title == source.lesson.content.title
    text = pdf_text(result)
    lesson = source.lesson.content
    assert lesson.title in text
    assert lesson.summary[0] in text
    assert f"Page 1 of {result.page_count}" in text
    assert "Rust lecture notes, pp. 3 to 4" in text  # source references
    assert "no longer available" in text  # a deleted document is said to be gone
    assert result.unsupported_characters == 0
    assert FILE_NAME.match(result.file_name)
    assert {"NotoSans", "NotoSans-Bold"} & embedded_fonts(result)


def test_the_notes_have_a_title_page_every_section_and_a_code_block() -> None:
    source = lesson_source()
    result = render("notes_pdf", source)
    text = pdf_text(result)
    reader = PdfReader(io.BytesIO(result.data))

    assert "STUDY NOTES" in reader.pages[0].extract_text()
    for section in source.lesson.content.sections:
        assert section.heading in text
    assert "input -> steps -> result" in text  # the sample lesson's code block
    assert "NotoSansMono" in embedded_fonts(result)
    # Headings are bookmarks too.
    assert [entry.title for entry in reader.outline if not isinstance(entry, list)][:1] == [
        "What you will learn"
    ]


def test_the_revision_sheet_is_short() -> None:
    notes = render("notes_pdf", lesson_source())
    sheet = render("revision_pdf", lesson_source())
    assert sheet.page_count < notes.page_count
    assert "Check yourself" in pdf_text(sheet)


def test_a_checklist_box_never_stays_behind_when_its_text_moves_to_the_next_page(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf = StudyDocument(_info(lesson_source(), "Revision sheet"))  # pyright: ignore[reportPrivateUsage]
    pdf.add_page()
    pdf.set_font(BODY_FAMILY, "", 9.5)
    pdf.set_y(pdf.page_break_trigger - 2)  # less room than one line
    boxes: list[int] = []
    draw = pdf.rect

    def rect(*args: Any, **kwargs: Any) -> Any:
        boxes.append(pdf.page)
        return draw(*args, **kwargs)

    monkeypatch.setattr(pdf, "rect", rect)
    _checklist(pdf, ["Explain what ownership means"])  # pyright: ignore[reportPrivateUsage]

    assert pdf.page == 2
    assert boxes == [2]


def test_a_lesson_without_citations_says_it_was_not_based_on_the_learners_documents() -> None:
    text = pdf_text(render("notes_pdf", lesson_source()))
    assert "does not cite any of your documents" in text


def test_sample_content_is_labelled_on_every_page() -> None:
    result = render("notes_pdf", lesson_source(is_synthetic=True))
    reader = PdfReader(io.BytesIO(result.data))
    for page in reader.pages:
        assert "Sample content from the test provider" in page.extract_text()


def test_a_roadmap_pdf_has_its_schedule_modules_and_open_questions() -> None:
    source = roadmap_source()
    result = render("roadmap_pdf", source)
    text = pdf_text(result)

    assert result.content_type == PDF_TYPE and result.page_count >= 2
    assert source.roadmap.title in text
    assert "Schedule at a glance" in text
    for module in source.roadmap.modules:
        assert module.title in text
    assert "Revision: Review the first week" in text
    assert "Optional extras" in text and "Rustonomicon" in text
    assert "Whether your course expects async Rust." in text
    assert FILE_NAME.match(result.file_name)


def test_a_long_roadmap_breaks_across_pages_with_the_table_header_repeated() -> None:
    result = render("roadmap_pdf", roadmap_source(topics=60))
    reader = PdfReader(io.BytesIO(result.data))
    headed = [p for p in reader.pages if "What to study" in p.extract_text()]
    assert len(headed) >= 2  # the schedule's header row repeats on the next page


# --- Unicode ------------------------------------------------------------------------------------


def test_hindi_and_odia_are_set_in_their_own_embedded_fonts() -> None:
    result = render("notes_pdf", lesson_source(content=indic_lesson(), language="hi"))

    fonts = embedded_fonts(result)
    assert {"NotoSansDevanagari", "NotoSansOriya"} <= {f.replace("-Bold", "") for f in fonts}
    assert result.unsupported_characters == 0
    # Every character of the Indic sample is covered by one of the bundled fonts.
    covered = coverage(("NotoSans", "NotoSansDevanagari", "NotoSansOriya"))
    assert all(ord(c) in covered for c in HINDI_TEXT + ODIA_TEXT if not c.isspace())
    # The language is declared for screen readers.
    reader = PdfReader(io.BytesIO(result.data))
    root: Any = reader.trailer["/Root"]
    assert root["/Lang"] == "hi"
    # The ASCII-only download name falls back to a generic one.
    assert result.file_name == "studypilot-notes.pdf"


def test_an_indic_page_is_actually_drawn() -> None:
    """Draw page 2 (where the Hindi and Odia paragraphs are) and check its text block has ink."""
    result = render("notes_pdf", lesson_source(content=indic_lesson(), language="hi"))
    document: Any = pypdfium2.PdfDocument(result.data)
    try:
        image: Any = document[1].render(scale=1).to_pil().convert("L")
    finally:
        document.close()
    width, height = image.size
    top: Any = image.crop((0, int(height * 0.15), width, int(height * 0.35)))
    histogram: list[int] = top.histogram()
    dark = sum(histogram[:128])
    assert dark > 500  # glyphs were drawn, not empty boxes or nothing


def test_characters_no_font_covers_are_replaced_visibly_and_counted() -> None:
    unsupported = "\U0001f600 த"  # an emoji and a Tamil letter: no bundled font has them

    def edit(data: dict[str, Any]) -> None:
        data["summary"][0] = f"Covered text then {unsupported} end."

    result = render("notes_pdf", lesson_source(content=lesson_with(edit)))
    assert result.unsupported_characters == 2
    assert REPLACEMENT in pdf_text(result)


def test_the_text_guard_keeps_joiners_and_drops_control_characters() -> None:
    guard = TextGuard()
    assert guard.clean("क्\u200dष") == "क्\u200dष"
    assert guard.clean("a\x07b\u202ec") == "abc"
    assert guard.unsupported == 0


# --- PowerPoint --------------------------------------------------------------------------------


def open_deck(result: RenderedFile) -> Any:
    return Presentation(io.BytesIO(result.data))


def text_of(slide: Any, name: str) -> str:
    """The text of the shape the writer named ``name`` ("Title", "Subtitle", "Body", ...)."""
    return next(shape.text_frame.text for shape in slide.shapes if shape.name == name)


def test_lesson_slides_are_a_valid_presentation_with_a_sensible_outline() -> None:
    source = lesson_source(citations=cited_sources())
    result = render("slides_pptx", source)

    assert result.content_type == PPTX_TYPE and FILE_NAME.match(result.file_name)
    with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
        names = set(archive.namelist())
        assert {"[Content_Types].xml", "ppt/presentation.xml"} <= names
        assert archive.testzip() is None
    deck = open_deck(result)
    assert len(deck.slides) == result.page_count
    assert deck.slide_width > deck.slide_height  # 16:9
    titles = [text_of(slide, "Title") for slide in deck.slides]
    assert titles[0] == source.lesson.content.title
    assert "What you will learn" in titles and "Summary" in titles and "Sources" in titles
    for section in source.lesson.content.sections:
        assert any(t.startswith(section.heading) for t in titles)


def test_no_slide_is_overcrowded() -> None:
    def edit(data: dict[str, Any]) -> None:
        long_paragraph = " ".join(["This sentence is deliberately long and wordy."] * 30)
        data["sections"][0]["blocks"] = [{"type": "paragraph", "text": long_paragraph}] * 9
        data["summary"] = [f"Summary point {n} " + "words " * 60 for n in range(8)]

    result = render("slides_pptx", lesson_source(content=lesson_with(edit)))
    outline = lesson_outline(lesson_source(content=lesson_with(edit)))
    for slide in outline:
        assert len(slide.points) <= max(MAX_POINTS, 6)
        assert all(len(point) <= MAX_POINT_CHARS for point in slide.points)
    assert any(slide.title.endswith("(continued)") for slide in outline)
    # The full paragraph text is kept in the speaker notes, not dropped.
    deck = open_deck(result)
    notes = [s.notes_slide.notes_text_frame.text for s in deck.slides if s.has_notes_slide]
    assert any("deliberately long and wordy" in n and len(n) > 1000 for n in notes)


def test_code_gets_its_own_slide_and_is_shortened_with_a_note_when_long() -> None:
    long_code = "\n".join(
        f"let value_{n} = compute_something_long(argument_{n});" for n in range(40)
    )

    def edit(data: dict[str, Any]) -> None:
        data["sections"][0]["blocks"].append(
            {"type": "code", "language": "rust", "code": long_code, "caption": "Many lines"}
        )

    outline = lesson_outline(lesson_source(content=lesson_with(edit)))
    code_slides = [s for s in outline if s.kind == "code"]
    long_one = next(s for s in code_slides if s.code and "value_0" in s.code)
    assert long_one.code is not None and len(long_one.code.splitlines()) == MAX_CODE_LINES
    assert long_one.caption is not None and "shortened" in long_one.caption
    assert long_one.notes == long_code


def test_a_list_item_cut_short_on_a_slide_is_kept_in_full_in_the_notes() -> None:
    long_item = (
        "A borrowed reference must never outlive the value it points to, which is exactly what "
        "the borrow checker proves before the program is allowed to compile and run"
    )
    assert len(long_item) > 120

    def edit(data: dict[str, Any]) -> None:
        data["sections"][0]["blocks"].append(
            {"type": "list", "ordered": False, "items": [long_item, "Short point"]}
        )

    source = lesson_source(content=lesson_with(edit))
    heading = source.lesson.content.sections[0].heading
    first = next(s for s in lesson_outline(source) if s.title == heading)
    assert any(point.endswith("…") and point[:-1] in long_item for point in first.points)
    assert first.notes is not None
    assert f"- {long_item}" in first.notes and "- Short point" in first.notes


def test_roadmap_module_notes_say_what_each_topic_is_and_what_to_practise() -> None:
    source = roadmap_source()
    module = source.roadmap.modules[0]
    slide = next(s for s in roadmap_outline(source) if s.title.startswith("Module 1:"))
    assert slide.notes is not None and slide.notes.startswith(module.summary)
    for topic in module.topics:
        assert f"{topic.title}: {topic.description}" in slide.notes
    assert "Practice: Write a short program that uses this idea." in slide.notes


def test_control_characters_in_code_are_dropped_not_written_out_as_escapes() -> None:
    code = "fn main() {\n\tlet bell = 1;\x07 // ring\x1b\n}\x00\n" + "\n".join(
        f"let value_{n} = {n};" for n in range(MAX_CODE_LINES + 4)
    )

    def edit(data: dict[str, Any]) -> None:
        data["sections"][0]["blocks"].append(
            {"type": "code", "language": "rust", "code": code, "caption": "Control characters"}
        )

    source = lesson_source(content=lesson_with(edit))
    slide = next(s for s in lesson_outline(source) if s.code and "bell" in s.code)
    for text in (slide.code, slide.notes):
        assert text is not None
        assert not any(char in text for char in "\x00\x07\x1b")
    assert slide.notes is not None and "\tlet bell = 1; // ring\n}" in slide.notes
    result = render("slides_pptx", source)
    with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
        xml = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith(("ppt/slides/slide", "ppt/notesSlides/"))
        )
    assert "_x0007_" not in xml and "_x001B_" not in xml and "_x0000_" not in xml
    assert "let bell = 1; // ring" in xml


def test_slides_name_a_complex_script_font_for_hindi_and_odia() -> None:
    result = render("slides_pptx", lesson_source(content=indic_lesson(), language="hi"))
    with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
        xml = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide")
        )
    assert '<a:cs typeface="Nirmala UI"/>' in xml
    assert 'lang="hi-IN"' in xml and 'lang="or-IN"' in xml
    assert HINDI_TEXT.split("।")[0] in xml


def test_control_characters_in_content_do_not_break_the_presentation() -> None:
    def edit(data: dict[str, Any]) -> None:
        data["sections"][0]["blocks"][0] = {"type": "paragraph", "text": "Bell\x07 and\x0b tab"}

    result = render("slides_pptx", lesson_source(content=lesson_with(edit)))
    assert open_deck(result).slides  # lxml would refuse these characters if they got through


def test_roadmap_slides_cover_every_module() -> None:
    source = roadmap_source(is_synthetic=True)
    result = render("slides_pptx", source)
    deck = open_deck(result)
    titles = [text_of(slide, "Title") for slide in deck.slides]
    for module in source.roadmap.modules:
        assert any(module.title in title for title in titles)
    assert "sample content" in text_of(deck.slides[0], "Subtitle")
    assert len(roadmap_outline(source)) == len(deck.slides)


# --- Small parts --------------------------------------------------------------------------------


def test_shorten_cuts_at_a_word_and_marks_the_cut() -> None:
    text = "word " * 100
    short = shorten(text, 40)
    assert len(short) <= 40 and short.endswith("…") and "wor…" not in short


def test_first_sentence_understands_the_devanagari_full_stop() -> None:
    assert first_sentence(HINDI_TEXT).endswith("।")
    assert first_sentence("One. Two.") == "One."


def test_plain_text_collapses_whitespace_and_drops_controls() -> None:
    assert plain("a\n\n b\x00\tc") == "a b c"


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Rust: Ownership & Borrowing!", "rust-ownership-borrowing-notes.pdf"),
        ("Café résumé", "cafe-resume-notes.pdf"),
        (HINDI_TEXT, "studypilot-notes.pdf"),
        ("x" * 300, "x" * 60 + "-notes.pdf"),
    ],
)
def test_download_names_are_safe(title: str, expected: str) -> None:
    name = file_name_for(title, "notes", "pdf")
    assert name == expected and FILE_NAME.match(name)


def test_a_format_that_does_not_fit_the_source_is_refused() -> None:
    with pytest.raises(UnsupportedExport):
        render("roadmap_pdf", lesson_source())
    with pytest.raises(UnsupportedExport):
        render("notes_pdf", roadmap_source())
